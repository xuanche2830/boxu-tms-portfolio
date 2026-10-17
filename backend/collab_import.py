"""Review first: source material is never an operational event or confirmed request."""
import base64
import binascii
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Literal

from fastapi import Depends, HTTPException, Request
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy import update
from sqlmodel import Session, select

from . import collab_auth as auth, collab_policy as policy, project_logistics as business
from .db import audit, identifier, now
from .import_models import ImportBatch
from .project_models import Project, ShippingRequest

MAX_FILE = 60_000
MAX_ROWS = 100
FIELDS = {
    'projects': {'code': '项目编号', 'name': '项目名称', 'owner': '负责人', 'department': '部门', 'note': '备注'},
    'requests': {'title': '需求名称', 'cargo': '货物', 'unit_kind': '单元类型', 'unit_count': '数量',
                 'origin': '起点', 'destination': '终点', 'requested_delivery': '期望到货日', 'owner': '负责人', 'note': '备注'},
}
KINDS = {'包裹': 'package', '托盘': 'pallet', '集装箱': 'container', '独立件': 'piece'}


class Upload(BaseModel):
    model_config = ConfigDict(extra='forbid')
    filename: str = Field(min_length=1, max_length=180)
    source_base64: str = Field(max_length=80_000)
    kind: Literal['projects', 'requests']
    project_id: str = Field(default='', max_length=80)


class DraftRow(BaseModel):
    model_config = ConfigDict(extra='forbid')
    row_number: int
    include: bool
    reviewed: bool
    fields: dict[str, str]


class SaveDraft(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1, strict=True)
    rows: list[DraftRow] = Field(min_length=1, max_length=MAX_ROWS)


class Confirm(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1, strict=True)
    confirm: Literal[True]


class ColumnMapping(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1, strict=True)
    mapping: dict[str, str]


def check_target(session, kind, project_id):
    policy.internal_only(session)
    if kind == 'projects':
        policy.owner_only(session)
        if project_id:
            raise HTTPException(422, '项目导入不需要指定目标项目')
    else:
        policy.project_access(session, project_id)


def owned_batch(session, key):
    policy.internal_only(session)
    batch = session.get(ImportBatch, key)
    if not batch or batch.created_by != policy.user(session).id:
        raise HTTPException(404, '导入批次不存在或没有访问权限')
    check_target(session, batch.kind, batch.project_id)
    return batch


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def parse_file(filename, content):
    suffix = Path(filename).suffix.lower()
    rows = []
    sheet = ''
    try:
        if suffix == '.csv':
            try:
                decoded = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                decoded = content.decode('gb18030')
            reader = csv.reader(io.StringIO(decoded), strict=True)
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError('最多导入100行，请拆分文件')
        elif suffix == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 200 or sum(i.file_size for i in entries) > 4_000_000:
                    raise ValueError('工作簿展开后过大，请只保留需要导入的工作表')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise ValueError('请提供仅含一张工作表的文件，避免遗漏其他工作表')
                ws = book.worksheets[0]
                sheet = ws.title
                if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > 25:
                    raise ValueError('最多100行、25列，请清除多余格式或拆分文件')
                for row in ws.iter_rows():
                    rows.append([cell_text(c.value) for c in row])
                    if len(rows) > MAX_ROWS + 1 or len(row) > 25:
                        raise ValueError('最多100行、25列')
            finally:
                book.close()
        else:
            raise ValueError('请选择 CSV 或 XLSX 文件')
    except HTTPException:
        raise
    except Exception as error:
        message = str(error) if isinstance(error, ValueError) else '无法解析文件，请另存为 CSV 或标准 XLSX 后重试'
        raise HTTPException(422, message)
    if len(rows) < 2 or len(rows[0]) > 25:
        raise HTTPException(422, '需要一行表头和至少一行数据，最多25列')
    headers = [cell_text(c).strip() for c in rows[0]]
    if any(not c for c in headers) or len(headers) != len(set(headers)):
        raise HTTPException(422, '表头不能留空或重复，请修改表头后重试')
    output = []
    for number, row in enumerate(rows[1:], 2):
        if not any(cell_text(c).strip() for c in row):
            continue
        if len(row) > len(headers):
            raise HTTPException(422, f'第{number}行比表头多出字段，请核对分隔符')
        raw = {h: cell_text(row[i]) if i < len(row) else '' for i, h in enumerate(headers)}
        if any(len(v) > 4000 for v in raw.values()):
            raise HTTPException(422, f'第{number}行有单元格超过4000字，请拆分内容')
        output.append({'row_number': number, 'sheet': sheet, 'raw': raw})
    if not output:
        raise HTTPException(422, '文件没有可预览的数据行')
    return output


def payload_for(batch, fields, number):
    data = {k: fields.get(k, '').strip() for k in FIELDS[batch.kind]}
    data['command_id'] = f'{batch.id}:{number}'
    if batch.kind == 'requests':
        data['project_id'] = batch.project_id
        count = data['unit_count']
        if re.fullmatch(r'[0-9]+', count):
            data['unit_count'] = int(count)
        data['unit_kind'] = KINDS.get(data['unit_kind'], data['unit_kind'])
    model = business.ProjectInput if batch.kind == 'projects' else business.RequestInput
    return model.model_validate(data)


def review(session, batch):
    reviewed = []
    seen = set()
    existing_requests = set()
    if batch.kind == 'requests' and batch.status == 'pending':
        for request in session.exec(select(ShippingRequest).where(ShippingRequest.project_id == batch.project_id)).all():
            existing_requests.add(json.dumps(request.model_dump(exclude={'id', 'version', 'created_at'}), sort_keys=True))
    for source in batch.rows:
        row = dict(source)
        errors = []
        try:
            payload = payload_for(batch, row['fields'], row['row_number'])
            signature = payload.code.upper() if batch.kind == 'projects' else json.dumps(payload.model_dump(exclude={'command_id'}), sort_keys=True)
            if signature in existing_requests:
                errors.append('目标项目已有相同需求，请排除此行；确为另一批发运时请明确修改需求名称')
            if row['include']:
                if signature in seen:
                    errors.append('本批次存在相同记录，请排除重复行或修正内容')
                seen.add(signature)
            if batch.kind == 'projects' and batch.status == 'pending' and session.exec(select(Project.id).where(Project.code == payload.code.upper())).first():
                errors.append('项目编号已存在；请排除此行，导入需求时选择已有项目')
        except ValidationError as error:
            for item in error.errors(include_input=False, include_url=False):
                key = str(item['loc'][0])
                hint = {'unit_count': '填写1至200的整数', 'requested_delivery': '填写有效日期 YYYY-MM-DD',
                        'unit_kind': '选择包裹、托盘、集装箱或独立件', 'code': '使用英文、数字、点、下划线或短横线'}
                errors.append(FIELDS[batch.kind].get(key, key) + '：' + hint.get(key, '必填，且不能超过字段长度限制'))
        row['errors'] = errors
        reviewed.append(row)
    return reviewed


def public_batch(session, batch):
    # SQLAlchemy expires instances on commit; refresh before serializing SQLModel.
    session.refresh(batch)
    return {**batch.model_dump(exclude={'source_base64', 'rows'}), 'rows': review(session, batch), 'field_labels': FIELDS[batch.kind]}


def register(app, signed_in):
    prefix = '/api/collab/imports'

    @app.get(prefix)
    def listing(session: Session = Depends(signed_in)):
        policy.internal_only(session)
        rows = session.exec(select(ImportBatch).where(ImportBatch.created_by == policy.user(session).id).order_by(ImportBatch.created_at.desc()).limit(100)).all()
        allowed = policy.grants(session, 'project') if policy.user(session).role != 'owner' else None
        return {'batches': [b.model_dump(exclude={'source_base64', 'rows', 'result'}) for b in rows
                            if allowed is None or b.kind == 'requests' and b.project_id in allowed]}

    @app.post(prefix)
    def upload(data: Upload, session: Session = Depends(signed_in)):
        check_target(session, data.kind, data.project_id)
        try:
            content = base64.b64decode(data.source_base64, validate=True)
        except (ValueError, binascii.Error):
            raise HTTPException(422, '文件编码无效，请重新选择文件')
        if not content or len(content) > MAX_FILE:
            raise HTTPException(422, '文件需要在1至60000字节之间，请拆分后导入')
        sha = hashlib.sha256(content).hexdigest()
        existing = session.exec(select(ImportBatch).where(ImportBatch.created_by == policy.user(session).id,
            ImportBatch.source_sha256 == sha, ImportBatch.kind == data.kind, ImportBatch.project_id == data.project_id)).first()
        if existing:
            return {**public_batch(session, existing), 'duplicate_source': True}
        rows = parse_file(data.filename, content)
        aliases = {alias: field for field, label in FIELDS[data.kind].items() for alias in [field, label]}
        for row in rows:
            fields = {key: '' for key in FIELDS[data.kind]}
            warnings = []
            assigned = set()
            for header, value in row['raw'].items():
                field = aliases.get(header)
                if field and field != 'note':
                    if field in assigned:
                        fields[field] = ''
                        warnings.append(f'{header}与其他列映射到同一字段，请手工确认')
                    else:
                        fields[field] = value.strip() if not value.startswith('=') else ''
                    assigned.add(field)
                elif value:
                    warnings.append(f'「{header}」只保留在原始来源，未写入业务字段')
                if value.startswith(('=', '+', '-', '@')):
                    warnings.append(f'「{header}」含公式或特殊开头，按原文保留，不执行')
            row.update(fields=fields, warnings=warnings, include=True, reviewed=False)
        batch = ImportBatch(id=identifier('IMP'), created_by=policy.user(session).id, filename=Path(data.filename.replace('\\', '/')).name,
            source_sha256=sha, source_base64=data.source_base64, kind=data.kind, project_id=data.project_id, created_at=now(), rows=rows)
        session.add(batch)
        auth.log(session, policy.user(session).id, 'import_preview', {'batch_id': batch.id, 'rows': len(rows), 'sha256': sha})
        session.commit()
        return public_batch(session, batch)

    @app.get(prefix + '/{key}')
    def detail(key: str, session: Session = Depends(signed_in)):
        return public_batch(session, owned_batch(session, key))

    @app.get(prefix + '/{key}/source')
    def source(key: str, session: Session = Depends(signed_in)):
        batch = owned_batch(session, key)
        return {'filename': batch.filename, 'sha256': batch.source_sha256, 'source_base64': batch.source_base64}

    @app.put(prefix + '/{key}/mapping')
    def mapping(key: str, data: ColumnMapping, session: Session = Depends(signed_in)):
        batch = owned_batch(session, key)
        if batch.status != 'pending':
            raise HTTPException(409, '已确认批次不能重新匹配字段')
        headers = set(batch.rows[0]['raw'])
        if set(data.mapping) != set(FIELDS[batch.kind]) or any(v and v not in headers for v in data.mapping.values()):
            raise HTTPException(422, '字段匹配必须选择原文件表头或留空')
        business.bump(session, batch, data.version)
        rows = []
        for row in batch.rows:
            fields = {key: row['raw'].get(header, '').strip() for key, header in data.mapping.items()}
            fields = {key: '' if value.startswith('=') else value for key, value in fields.items()}
            rows.append({**row, 'fields': fields, 'reviewed': False, 'mapping': data.mapping})
        batch.rows = rows
        session.add(batch)
        auth.log(session, policy.user(session).id, 'import_columns_mapped', {'batch_id': key, 'mapping': data.mapping})
        session.commit()
        return public_batch(session, batch)

    @app.put(prefix + '/{key}')
    def save(key: str, data: SaveDraft, session: Session = Depends(signed_in)):
        batch = owned_batch(session, key)
        if batch.status != 'pending':
            raise HTTPException(409, '批次已确认，原始来源与确认内容不能再修改')
        if {r.row_number for r in data.rows} != {r['row_number'] for r in batch.rows} or len(data.rows) != len(batch.rows):
            raise HTTPException(422, '草稿行号必须与原始文件一致；跳过记录请使用排除')
        for row in data.rows:
            if set(row.fields) != set(FIELDS[batch.kind]) or any(len(v) > 2000 for v in row.fields.values()):
                raise HTTPException(422, '草稿字段不完整或内容过长')
        business.bump(session, batch, data.version)
        edits = {r.row_number: r for r in data.rows}
        batch.rows = [{**r, **edits[r['row_number']].model_dump()} for r in batch.rows]
        session.add(batch)
        auth.log(session, policy.user(session).id, 'import_draft_saved', {'batch_id': key, 'version': batch.version})
        session.commit()
        return public_batch(session, batch)

    @app.post(prefix + '/{key}/confirm')
    def confirm(key: str, data: Confirm, request: Request, session: Session = Depends(signed_in)):
        batch = owned_batch(session, key)
        # A no-op UPDATE obtains a write lock on SQLite and a row lock on PostgreSQL.
        # Serialize both retries of one batch and different batches for one project.
        # Do not increment the public project version just to acquire this lock.
        session.exec(update(ImportBatch).where(ImportBatch.id == key).values(version=ImportBatch.version))
        if batch.kind == 'requests':
            session.exec(update(Project).where(Project.id == batch.project_id).values(version=Project.version))
        session.expire_all()
        auth.authenticate(session, request)
        batch = owned_batch(session, key)
        if batch.status == 'confirmed':
            return {**public_batch(session, batch), 'replayed': True}
        if batch.version != data.version:
            raise HTTPException(409, '草稿已更新，请重新打开批次核对后再确认')
        selected = [r for r in review(session, batch) if r['include']]
        if not selected or any(r['errors'] or not r['reviewed'] for r in selected):
            raise HTTPException(422, '请修正错误，并逐行标记已核对；至少保留一行')
        business.bump(session, batch, data.version)
        results = []
        for row in selected:
            payload = payload_for(batch, row['fields'], row['row_number'])
            handler = business.create_project if batch.kind == 'projects' else business.create_request
            result = handler(session, payload)
            try:
                session.flush()
            except IntegrityError:
                session.rollback()
                raise HTTPException(409, '导入内容与已有记录冲突，整批尚未写入；请重新打开批次核对')
            lineage = {**result, 'row_number': row['row_number'], 'batch_id': batch.id, 'source_sha256': batch.source_sha256}
            audit(session, result['id'], 'import_confirmed', None, lineage)
            results.append(lineage)
        batch.status, batch.confirmed_at, batch.result = 'confirmed', now(), results
        session.add(batch)
        auth.log(session, policy.user(session).id, 'import_confirmed', {'batch_id': key, 'records': len(results)})
        session.commit()
        return public_batch(session, batch)
